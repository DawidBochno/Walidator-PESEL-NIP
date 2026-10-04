#!/usr/bin/env python3
"""Walidator PESEL / NIP / REGON / IBAN / nr dowodu w arkuszach Excela.
Bledne komorki podswietla na czerwono z komentarzem (powod bledu),
z PESEL dopisuje date urodzenia i plec. Wynik: kopia pliku z _sprawdzony.

Uruchomienie: python walidator.py            (GUI)
              python walidator.py --selftest (test logiki)
"""
import datetime
import os
import re
import sys
import threading
import traceback

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

EXTS = (".xlsx", ".xlsm")
KINDS = ["PESEL", "NIP", "REGON", "IBAN", "Dowód"]
# naglowek kolumny -> rodzaj numeru (pierwsze trafienie wygrywa)
HEADER_RX = [
    (r"pesel", "PESEL"),
    (r"regon", "REGON"),
    (r"\bnip\b", "NIP"),
    (r"iban|rachun|nr konta|numer konta", "IBAN"),
    (r"dow[oó]d", "Dowód"),
]
# log bez polskich znakow (konsola cp1250 / cp1252 w CI)
ASCII = str.maketrans("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ", "acelnoszzACELNOSZZ")
CENTURY = {0: 1900, 1: 2000, 2: 2100, 3: 2200, 4: 1800}

# ------------------------------------------------------------ walidatory ----


def _weighted(s, weights):
    return sum(int(a) * b for a, b in zip(s, weights))


def pesel_info(s):
    """(data urodzenia, 'K'/'M') albo None gdy data jest niemozliwa."""
    m = int(s[2:4])
    try:
        born = datetime.date(CENTURY[m // 20] + int(s[:2]), m % 20, int(s[4:6]))
    except ValueError:
        return None
    return born, "K" if int(s[9]) % 2 == 0 else "M"


def iban_ok(s):
    return int("".join(str(int(c, 36)) for c in s[4:] + s[:4])) % 97 == 1


def normalize(value, kind):
    """Wartosc komorki -> tekst do sprawdzenia. Liczby z Excela traca zera
    z przodu (PESEL 02... zapisany jako liczba), wiec sa uzupelniane."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and not isinstance(value, bool):
        s = str(value)
        width = {"PESEL": 11, "NIP": 10, "REGON": 9 if len(s) <= 9 else 14}.get(kind)
        return s.zfill(width) if width else s
    s = re.sub(r"[\s\-.]", "", str(value)).upper()
    if kind == "NIP" and s.startswith("PL"):
        s = s[2:]
    if kind == "IBAN" and s[:1].isdigit():
        s = "PL" + s
    return s


def check(kind, value):
    """Zwraca powod bledu albo None, gdy numer jest poprawny."""
    s = normalize(value, kind)
    if kind == "PESEL":
        if not re.fullmatch(r"\d{11}", s):
            return "PESEL musi mieć 11 cyfr"
        if (10 - _weighted(s[:10], [1, 3, 7, 9] * 3) % 10) % 10 != int(s[10]):
            return "błędna cyfra kontrolna"
        if pesel_info(s) is None:
            return "niemożliwa data urodzenia"
    elif kind == "NIP":
        if not re.fullmatch(r"\d{10}", s):
            return "NIP musi mieć 10 cyfr"
        if _weighted(s, [6, 5, 7, 2, 3, 4, 5, 6, 7]) % 11 != int(s[9]):
            return "błędna cyfra kontrolna"
    elif kind == "REGON":
        if not re.fullmatch(r"\d{9}|\d{14}", s):
            return "REGON musi mieć 9 albo 14 cyfr"
        w = [8, 9, 2, 3, 4, 5, 6, 7] if len(s) == 9 else [2, 4, 8, 5, 0, 9, 7, 3, 6, 1, 2, 4, 8]
        if _weighted(s, w) % 11 % 10 != int(s[-1]):
            return "błędna cyfra kontrolna"
    elif kind == "IBAN":
        if not re.fullmatch(r"[A-Z]{2}\d{2}[0-9A-Z]{11,30}", s):
            return "nieprawidłowy format numeru konta"
        if s.startswith("PL") and len(s) != 28:
            return "polski numer konta musi mieć 26 cyfr"
        if not iban_ok(s):
            return "błędna suma kontrolna (literówka w numerze konta)"
    elif kind == "Dowód":
        if not re.fullmatch(r"[A-Z]{3}\d{6}", s):
            return "format: 3 litery i 6 cyfr"
        vals = [ord(c) - 55 if c.isalpha() else int(c) for c in s]
        if sum(a * b for a, b in zip(vals, [7, 3, 1, 9, 7, 3, 1, 7, 3])) % 10:
            return "błędna cyfra kontrolna"
    return None


# ------------------------------------------------------------- kolumny ----


def kind_by_name(name):
    name = name.strip().lower()
    for k in KINDS:
        if k.lower() == name or (k == "Dowód" and name == "dowod"):
            return k
    return None


def parse_mapping(text):
    """'C=PESEL, F=nip' -> {3: 'PESEL', 6: 'NIP'}. Bledny wpis -> ValueError."""
    from openpyxl.utils import column_index_from_string
    out = {}
    for part in re.split(r"[,;\n]+", text):
        if not part.strip():
            continue
        m = re.fullmatch(r"\s*([A-Za-z]{1,3})\s*=\s*(\S+)\s*", part)
        kind = m and kind_by_name(m.group(2))
        if not kind:
            raise ValueError("Nie rozumiem: %r (wpisz np. C=PESEL, D=NIP)" % part.strip())
        out[column_index_from_string(m.group(1).upper())] = kind
    return out


def detect_columns(ws):
    """Rodzaj numeru po naglowku w wierszu 1: {numer_kolumny: rodzaj}."""
    out = {}
    for cell in ws[1]:
        h = str(cell.value or "").lower()
        for rx, kind in HEADER_RX:
            if re.search(rx, h):
                out[cell.column] = kind
                break
    return out


# ---------------------------------------------------------------- plik ----


def check_file(src, dst, mapping=None, log=print):
    """Sprawdza wszystkie arkusze. Zwraca {rodzaj: [ok, bledne]}."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    red = PatternFill("solid", fgColor="FFC7CE")
    wb = openpyxl.load_workbook(src, keep_vba=src.lower().endswith(".xlsm"))
    stats = {}
    errors = []
    for ws in wb.worksheets:
        cols = mapping or detect_columns(ws)
        if not cols:
            continue
        found = ", ".join("%s=%s" % (get_column_letter(c), k) for c, k in sorted(cols.items()))
        log(("  Arkusz '%s': %s" % (ws.title, found)).translate(ASCII))
        extra = {}  # kolumna PESEL -> pierwsza z dwoch dopisanych kolumn
        for c in sorted(c for c, k in cols.items() if k == "PESEL"):
            first = ws.max_column + 1
            extra[c] = first
            head = str(ws.cell(1, c).value or "PESEL")
            for i, label in enumerate(("data urodzenia", "płeć")):
                cell = ws.cell(1, first + i, "%s – %s" % (head, label))
                cell.font = Font(bold=True)
                ws.column_dimensions[get_column_letter(first + i)].width = 16
        for row in range(2, ws.max_row + 1):
            for c, kind in cols.items():
                cell = ws.cell(row, c)
                if cell.value is None or str(cell.value).strip() == "":
                    continue
                st = stats.setdefault(kind, [0, 0])
                err = check(kind, cell.value)
                if err:
                    st[1] += 1
                    cell.fill = red
                    cell.comment = Comment("%s: %s" % (kind, err), "Walidator")
                    errors.append("%s!%s  %s  %s" % (ws.title, cell.coordinate, cell.value, err))
                    continue
                st[0] += 1
                if c in extra:
                    born, sex = pesel_info(normalize(cell.value, kind))
                    ws.cell(row, extra[c], born).number_format = "yyyy-mm-dd"
                    ws.cell(row, extra[c] + 1, sex)
    if not stats:
        log("  Brak numerow do sprawdzenia (kolumny rozpoznawane po naglowku w wierszu 1)"
            " - mozna je wpisac recznie, np. C=PESEL. Plik nie zapisany.")
        return stats
    for e in errors[:20]:
        log(("    BLAD " + e).translate(ASCII))
    if len(errors) > 20:
        log("    ... i jeszcze %d (zaznaczone na czerwono w pliku)" % (len(errors) - 20))
    wb.save(dst)
    return stats


def run(inp, out_dir, mapping_text="", log=print):
    mapping = parse_mapping(mapping_text)
    if os.path.isfile(inp):
        files = [inp]
    else:
        files = sorted(os.path.join(inp, f) for f in os.listdir(inp)
                       if f.lower().endswith(EXTS) and not f.startswith("~$"))
    if not files:
        log("Brak plikow XLSX w: %s" % inp)
        return
    os.makedirs(out_dir, exist_ok=True)
    for f in files:
        name, ext = os.path.splitext(os.path.basename(f))
        dst = os.path.join(out_dir, name + "_sprawdzony" + ext.lower())
        log("Sprawdzam: %s" % os.path.basename(f))
        try:
            stats = check_file(f, dst, mapping, log)
            if stats:
                log("  " + ", ".join("%s: %d OK, %d bledne" % (k, ok, bad)
                                     for k, (ok, bad) in stats.items()).translate(ASCII))
                log("  -> %s" % dst)
        except Exception:
            log("BLAD: %s\n%s" % (os.path.basename(f), traceback.format_exc()))
    log("Zakonczono (%d plikow)." % len(files))


# -------------------------------------------------------------------- GUI ----


def gui():
    import tkinter as tk
    from tkinter import filedialog, ttk, scrolledtext

    root = tk.Tk()
    root.title("Walidator PESEL / NIP / REGON / IBAN")
    root.geometry("780x520")
    pad = dict(padx=6, pady=3)

    v_in = tk.StringVar(value=os.path.join(APP_DIR, "INPUT"))
    v_out = tk.StringVar(value=os.path.join(APP_DIR, "OUTPUT"))
    v_map = tk.StringVar()

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    ttk.Label(f, text="Plik lub folder:").grid(row=0, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_in, width=60).grid(row=0, column=1, **pad)
    ttk.Button(f, text="Plik...", command=lambda: v_in.set(filedialog.askopenfilename(
        filetypes=[("Excel", "*.xlsx *.xlsm")]) or v_in.get())).grid(row=0, column=2, **pad)
    ttk.Button(f, text="Folder...", command=lambda: v_in.set(
        filedialog.askdirectory() or v_in.get())).grid(row=0, column=3, **pad)
    ttk.Label(f, text="Folder wyjsciowy:").grid(row=1, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_out, width=60).grid(row=1, column=1, **pad)
    ttk.Button(f, text="Wybierz...", command=lambda: v_out.set(
        filedialog.askdirectory() or v_out.get())).grid(row=1, column=2, **pad)
    ttk.Label(f, text="Kolumny:").grid(row=2, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_map, width=60).grid(row=2, column=1, **pad)
    ttk.Label(f, text="puste = rozpoznaj po naglowku; recznie np.  C=PESEL, D=NIP, E=IBAN",
              foreground="gray").grid(row=3, column=1, sticky="w", padx=6)

    log_box = scrolledtext.ScrolledText(root, height=16)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    btn = ttk.Button(root, text="Sprawdz")
    btn.pack(pady=6)
    log_box.pack(fill="both", expand=True, **pad)

    def start():
        inp, out = v_in.get().strip('" '), v_out.get().strip('" ')
        log_box.delete("1.0", "end")
        if not os.path.exists(inp):
            return log("Wskaz istniejacy plik lub folder.")
        if not out:
            return log("Wskaz folder wyjsciowy.")
        try:
            parse_mapping(v_map.get())
        except ValueError as e:
            return log(str(e))
        btn.config(state="disabled")

        def work():
            try:
                run(inp, out, v_map.get(), log)
            finally:
                root.after(0, lambda: btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Walidator-PESEL-NIP", "main", "walidator.py")
    root.mainloop()


# --------------------------------------------------------------- selftest ----


def selftest():
    import tempfile
    import openpyxl

    assert check("PESEL", "44051401359") is None
    assert check("PESEL", "44051401358") == "błędna cyfra kontrolna"
    assert check("PESEL", "4405140135") == "PESEL musi mieć 11 cyfr"
    assert check("PESEL", "90010100016") is None
    assert check("PESEL", "90023100014") == "niemożliwa data urodzenia"  # 31 lutego
    assert check("PESEL", 2270803624) is None  # liczba w Excelu, zgubione 0 z przodu
    assert pesel_info("02270803624") == (datetime.date(2002, 7, 8), "K")
    assert pesel_info("44051401359") == (datetime.date(1944, 5, 14), "M")
    assert check("NIP", "PL 123-456-32-18") is None and check("NIP", 1234563219)
    assert check("REGON", "123456785") is None and check("REGON", "12345678512347") is None
    assert check("REGON", "123456789") == "błędna cyfra kontrolna"
    assert check("IBAN", "PL61 1090 1014 0000 0712 1981 2874") is None
    assert check("IBAN", "61109010140000071219812874") is None
    assert check("IBAN", "61109010140000071219812875").startswith("błędna")
    assert check("IBAN", "DE89 3704 0044 0532 0130 00") is None
    assert check("Dowód", "aba 300000") is None and check("Dowód", "ABA300001")
    assert parse_mapping(" c=pesel; AA=Dowod ") == {3: "PESEL", 27: "Dowód"}
    try:
        parse_mapping("C-PESEL")
        raise AssertionError("parse_mapping przyjal bledny wpis")
    except ValueError:
        pass

    tmp = tempfile.mkdtemp()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append(["Nazwisko", "PESEL", "NIP firmy", "Nr rachunku", "Uwagi"])
    ws.append(["Kowalska", "02270803624", "123-456-32-18", "PL61 1090 1014 0000 0712 1981 2874", "x"])
    ws.append(["Nowak", "44051401358", "1234563219", "", None])
    ws.append(["Wiśniewski", 44051401359, None, "61109010140000071219812875", 12345])
    wb.create_sheet("Pusty").append(["Cos", "Innego"])
    wb.save(os.path.join(tmp, "lista.xlsx"))
    open(os.path.join(tmp, "~$lista.xlsx"), "wb").close()

    lines = []
    run(tmp, os.path.join(tmp, "out"), "", lines.append)
    assert not any(x.startswith("BLAD:") for x in lines), lines
    out = openpyxl.load_workbook(os.path.join(tmp, "out", "lista_sprawdzony.xlsx"))
    ws = out["Lista"]
    assert all(ord(ch) < 128 for line in lines for ch in line), lines
    assert any("B=PESEL, C=NIP, D=IBAN" in x for x in lines), lines
    bad = sorted(c.coordinate for row in ws.iter_rows() for c in row if c.comment)
    assert bad == ["B3", "C3", "D4"], bad
    assert ws["B3"].comment.text == "PESEL: błędna cyfra kontrolna"
    assert ws["B3"].fill.fgColor.rgb.endswith("FFC7CE") and not ws["B2"].fill.fill_type
    assert ws["F1"].value == "PESEL – data urodzenia" and ws["G1"].value == "PESEL – płeć"
    assert ws["F2"].value.date() == datetime.date(2002, 7, 8) and ws["G2"].value == "K"
    assert ws["F3"].value is None and ws["G4"].value == "M"
    assert any("PESEL: 2 OK, 1 bledne" in x for x in lines), lines

    # reczne kolumny: sprawdza tylko wskazane
    run(os.path.join(tmp, "lista.xlsx"), os.path.join(tmp, "out2"), "C=NIP", lines.append)
    ws = openpyxl.load_workbook(os.path.join(tmp, "out2", "lista_sprawdzony.xlsx"))["Lista"]
    assert [c.coordinate for row in ws.iter_rows() for c in row if c.comment] == ["C3"]
    assert ws["F1"].value is None

    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        if "--gui" in sys.argv:
            gui()
    else:
        gui()
